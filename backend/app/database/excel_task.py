from __future__ import annotations

import threading
import sqlite3
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "TASK_ID",
    "PROJECT_CODE",
    "PROJECT_NAME",
    "SEGMENT_NAME",
    "MACHINE_NO",
    "STATUS",
    "START_TIME",
    "END_TIME",
    "WORK_TIME",
    "OP_DATE",
    "OP_PROGRAM",
    "OP_LENGTH_CUT",
    "OP_LENGTH_GROOVE",
    "REMARK",
}


class ExcelTaskDataError(RuntimeError):
    pass


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def _decimal(value: Any) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _valid_work_time(row: dict[str, Any]) -> bool:
    try:
        return _decimal(row.get("WORK_TIME")) >= 0
    except (ArithmeticError, ValueError):
        return False


class ExcelTaskDataSource:
    _lock = threading.RLock()
    _cache: dict[Path, tuple[tuple[int, int], tuple[str, ...], tuple[dict[str, Any], ...]]] = {}
    _sql_cache: dict[Path, tuple[tuple[int, int], sqlite3.Connection]] = {}

    def __init__(self, path: Path) -> None:
        self.path = path.resolve()

    def _signature(self) -> tuple[int, int]:
        if not self.path.is_file():
            raise ExcelTaskDataError(f"Excel task 数据文件不存在：{self.path}")
        stat = self.path.stat()
        return stat.st_mtime_ns, stat.st_size

    def _load(self) -> tuple[tuple[str, ...], tuple[dict[str, Any], ...]]:
        signature = self._signature()
        with self._lock:
            cached = self._cache.get(self.path)
            if cached and cached[0] == signature:
                return cached[1], cached[2]

            workbook = load_workbook(self.path, read_only=True, data_only=True)
            try:
                if "TASK" not in workbook.sheetnames:
                    raise ExcelTaskDataError("TASK.xlsx 中缺少名为 TASK 的工作表")
                worksheet = workbook["TASK"]
                iterator = worksheet.iter_rows(values_only=True)
                raw_headers = next(iterator, None)
                if not raw_headers:
                    raise ExcelTaskDataError("TASK 工作表没有表头")
                headers = tuple(str(value).strip() if value is not None else "" for value in raw_headers)
                missing = sorted(REQUIRED_COLUMNS - set(headers))
                if missing:
                    raise ExcelTaskDataError(f"TASK 工作表缺少必要字段：{', '.join(missing)}")
                if len(set(headers)) != len(headers):
                    raise ExcelTaskDataError("TASK 工作表存在重复字段名")
                rows = tuple(
                    dict(zip(headers, values, strict=True))
                    for values in iterator
                    if any(value not in (None, "") for value in values)
                )
            finally:
                workbook.close()
            self._cache[self.path] = (signature, headers, rows)
            return headers, rows

    @staticmethod
    def production_date(row: dict[str, Any]) -> date | None:
        return _date_value(row.get("OP_DATE")) or _date_value(row.get("START_TIME"))

    def headers(self) -> list[str]:
        return list(self._load()[0])

    def all_rows(self) -> list[dict[str, Any]]:
        return list(self._load()[1])

    def completed_rows(
        self,
        *,
        start: date,
        end: date,
        project_code: str | None = None,
        ship_no: str | None = None,
        machine_no: str | None = None,
    ) -> list[dict[str, Any]]:
        result = []
        for row in self._load()[1]:
            try:
                status = int(row.get("STATUS"))
            except (TypeError, ValueError):
                continue
            production_date = self.production_date(row)
            if (
                status != 2
                or not _valid_work_time(row)
                or production_date is None
                or not (start <= production_date < end)
            ):
                continue
            project_name = _clean_text(row.get("PROJECT_NAME"))
            legacy_project_code = _clean_text(row.get("PROJECT_CODE"))
            if project_code and project_code not in {project_name, legacy_project_code}:
                continue
            if ship_no and ship_no != project_name:
                continue
            if machine_no and (
                _clean_text(row.get("MACHINE_NO")) or ""
            ).casefold() != machine_no.strip().casefold():
                continue
            result.append(row)
        return result

    def list_projects(self, machine_no: str | None = None) -> list[str]:
        return sorted(
            {
                value
                for row in self._load()[1]
                if not machine_no
                or (_clean_text(row.get("MACHINE_NO")) or "").casefold()
                == machine_no.strip().casefold()
                if (
                    value := _clean_text(row.get("PROJECT_NAME"))
                    or _clean_text(row.get("PROJECT_CODE"))
                )
            }
        )

    def latest_completed_date(
        self, on_or_before: date, machine_no: str | None = None
    ) -> date | None:
        latest: date | None = None
        for row in self._load()[1]:
            try:
                status = int(row.get("STATUS"))
            except (TypeError, ValueError):
                continue
            production_date = self.production_date(row)
            if (
                status == 2
                and _valid_work_time(row)
                and (
                    not machine_no
                    or (_clean_text(row.get("MACHINE_NO")) or "").casefold()
                    == machine_no.strip().casefold()
                )
                and production_date is not None
                and production_date <= on_or_before
                and (latest is None or production_date > latest)
            ):
                latest = production_date
        return latest

    @staticmethod
    def normalized(row: dict[str, Any]) -> dict[str, Any]:
        project_name = _clean_text(row.get("PROJECT_NAME"))
        return {
            "task_id": str(row.get("TASK_ID")),
            "project_code": _clean_text(row.get("PROJECT_CODE")),
            "project_name": project_name,
            "ship_no": None,
            "segment_name": _clean_text(row.get("SEGMENT_NAME")),
            "program_no": _clean_text(row.get("OP_PROGRAM")),
            "equipment_id": _clean_text(row.get("MACHINE_NO")),
            "cut_length": _decimal(row.get("OP_LENGTH_CUT")) / Decimal("1000"),
            "groove_length": _decimal(row.get("OP_LENGTH_GROOVE")) / Decimal("1000"),
            "work_time": _decimal(row.get("WORK_TIME")),
            "production_time": row.get("START_TIME"),
            "source_updated_at": row.get("END_TIME"),
            "remark": _clean_text(row.get("REMARK")),
        }

    def status(self) -> dict[str, Any]:
        _, rows = self._load()
        statuses = Counter(row.get("STATUS") for row in rows)
        completed_rows = [row for row in rows if row.get("STATUS") == 2]
        excluded_negative_work_time_count = sum(
            1 for row in completed_rows if not _valid_work_time(row)
        )
        machines = Counter(
            value for row in rows if (value := _clean_text(row.get("MACHINE_NO")))
        )
        dates = [value for row in rows if (value := self.production_date(row))]
        signature = self._signature()
        return {
            "path": str(self.path),
            "filename": self.path.name,
            "file_size": signature[1],
            "row_count": len(rows),
            "completed_count": int(statuses.get(2, 0)),
            "reportable_completed_count": len(completed_rows)
            - excluded_negative_work_time_count,
            "excluded_negative_work_time_count": excluded_negative_work_time_count,
            "date_min": min(dates).isoformat() if dates else None,
            "date_max": max(dates).isoformat() if dates else None,
            "machines": sorted(machines),
            "modified_at": datetime.fromtimestamp(self.path.stat().st_mtime).isoformat(timespec="seconds"),
        }

    @staticmethod
    def _sqlite_value(value: Any) -> Any:
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, Decimal):
            return float(value)
        return value

    def _sqlite_connection(self) -> sqlite3.Connection:
        signature = self._signature()
        with self._lock:
            cached = self._sql_cache.get(self.path)
            if cached and cached[0] == signature:
                return cached[1]
            if cached:
                cached[1].close()
            headers, rows = self._load()
            connection = sqlite3.connect(":memory:", check_same_thread=False)
            connection.row_factory = sqlite3.Row
            connection.execute("ATTACH DATABASE ':memory:' AS dbo")
            quoted_columns = ", ".join(f'"{column}"' for column in headers)
            connection.execute(f"CREATE TABLE dbo.task ({quoted_columns})")
            placeholders = ",".join("?" for _ in headers)
            connection.executemany(
                f"INSERT INTO dbo.task ({quoted_columns}) VALUES ({placeholders})",
                (
                    tuple(self._sqlite_value(row[column]) for column in headers)
                    for row in rows
                ),
            )
            connection.commit()
            self._sql_cache[self.path] = (signature, connection)
            return connection

    def execute_select(
        self,
        sql: str,
        parameters: tuple[Any, ...],
        max_rows: int,
    ) -> list[dict[str, Any]]:
        _, rows = self.execute_select_with_columns(sql, parameters, max_rows)
        return rows

    def execute_select_with_columns(
        self,
        sql: str,
        parameters: tuple[Any, ...],
        max_rows: int,
    ) -> tuple[list[str], list[dict[str, Any]]]:
        cleaned = sql.strip().rstrip(";")
        limited_sql = f"SELECT * FROM ({cleaned}) AS local_query_result LIMIT ?"
        values = tuple(self._sqlite_value(value) for value in parameters) + (max_rows,)
        with self._lock:
            cursor = self._sqlite_connection().execute(limited_sql, values)
            columns = [str(item[0]) for item in cursor.description or ()]
            return columns, [dict(row) for row in cursor.fetchall()]
